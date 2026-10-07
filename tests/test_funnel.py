import hashlib
import hmac
import io
import json
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import load_workbook

import crm
import export
import funnel
from bot import stats_text
from quiz import QUIZ

ANSWERS = {item["key"]: item["options"][0] for item in QUIZ}


@pytest.fixture(autouse=True)
def fresh_db(tmp_path, monkeypatch):
    monkeypatch.setattr(funnel, "PATH", str(tmp_path / "quiz.db"))


def person(user_id: int, username: str | None = "polina"):
    return SimpleNamespace(id=user_id, username=username, full_name="Полина")


def walk(user_id: int, steps: int, source: str = "reels") -> None:
    funnel.track(user_id, "begin", source)
    for step in range(steps):
        funnel.track(user_id, "answer", source, step)


def test_funnel_counts_drop_off_by_question():
    walk(1, len(QUIZ))
    walk(2, 3)
    walk(3, 1, source="ads")
    funnel.save_lead(person(1), ANSWERS, "reels")
    funnel.set_phone(1, "+79991112233")

    data = funnel.report(len(QUIZ))
    assert data["started"] == 3
    assert data["answered"][0] == 3
    assert data["answered"][2] == 2
    assert data["answered"][-1] == 1
    assert data["finished"] == 1 and data["phones"] == 1
    assert data["sources"] == [("reels", 1)]

    text = stats_text(data)
    assert "Начали: 3" in text
    assert "Ответили на 3-й: 2 (67%)" in text
    assert "Оставили телефон: 1 (100%)" in text


def test_stalled_quiz_is_reminded_once():
    funnel.progress(7, 70, 4, "reels")
    assert funnel.stalled(3600) == []
    later = funnel.stalled(3600, now=10**10)
    assert [(row["user_id"], row["step"]) for row in later] == [(7, 4)]
    funnel.reminded(7)
    assert funnel.stalled(3600, now=10**10) == []

    funnel.progress(7, 70, 5, "reels")
    assert funnel.stalled(3600, now=10**10)
    funnel.done(7)
    assert funnel.stalled(3600, now=10**10) == []


def test_phone_goes_to_the_latest_lead_and_triggers_sending():
    funnel.save_lead(person(5), ANSWERS, "—")
    assert funnel.unsent(600) == []
    lead = funnel.set_phone(5, "+79990001122")
    assert lead["phone"] == "+79990001122"
    [pending] = funnel.unsent(600)
    funnel.mark_sent(pending["id"])
    assert funnel.unsent(600, now=10**10) == []


def test_phone_without_quiz_is_ignored():
    assert funnel.set_phone(404, "+7999") is None


def test_export_has_every_answer_column():
    funnel.save_lead(person(9, username=None), ANSWERS, "ads")
    sheet = load_workbook(io.BytesIO(export.leads_xlsx(funnel.leads()))).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][5:] == tuple(item["q"] for item in QUIZ)
    assert rows[1][2] == 9 and rows[1][4] == "ads"


async def test_lead_is_signed_for_ai_crm(monkeypatch):
    calls = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(200, json={"id": 31, "status": "ok"})

    monkeypatch.setattr(crm, "URL", "https://crm.example.ru/webhook/lead")
    monkeypatch.setattr(crm, "SECRET", "shared")
    monkeypatch.setattr(crm, "TRANSPORT", httpx.MockTransport(handler))

    funnel.save_lead(person(3), ANSWERS, "reels")
    lead = funnel.set_phone(3, "+79991112233")
    assert await crm.send(lead) == 31

    [call] = calls
    expected = hmac.new(
        b"shared", call.headers["x-timestamp"].encode() + b"." + call.content, hashlib.sha256
    ).hexdigest()
    assert call.headers["x-signature"] == expected
    body = json.loads(call.content)
    assert body["contact"] == "+79991112233" and body["consent"] is True
    assert body["utm_source"] == "reels" and body["source"] == "квиз-бот"
    assert QUIZ[0]["q"] in body["text"]


async def test_crm_failure_keeps_the_lead_unsent(monkeypatch):
    monkeypatch.setattr(crm, "URL", "https://crm.example.ru/webhook/lead")
    monkeypatch.setattr(crm, "SECRET", "shared")
    monkeypatch.setattr(crm, "TRANSPORT", httpx.MockTransport(lambda r: httpx.Response(401)))
    lead = {"user_id": 1, "username": None, "answers": ANSWERS, "source": "—"}
    assert await crm.send(lead) is None
    assert crm.payload(lead)["contact"] == "tg id 1"
    assert crm.payload(lead)["consent"] is False


async def test_contact_handler_saves_phone_and_pushes_to_crm(monkeypatch):
    from unittest.mock import AsyncMock

    import bot

    sent = []

    async def fake_send(lead):
        sent.append(lead)
        return 77

    monkeypatch.setattr(crm, "URL", "https://crm.example.ru/webhook/lead")
    monkeypatch.setattr(crm, "SECRET", "shared")
    monkeypatch.setattr(crm, "send", fake_send)
    monkeypatch.setattr(bot, "ADMIN_CHAT_ID", "-100")

    funnel.save_lead(person(11), ANSWERS, "reels")
    message = SimpleNamespace(
        contact=SimpleNamespace(user_id=11, phone_number="+79995554433"),
        from_user=SimpleNamespace(id=11),
        answer=AsyncMock(),
        bot=SimpleNamespace(send_message=AsyncMock()),
    )
    await bot.on_contact(message)

    assert sent[0]["phone"] == "+79995554433"
    assert sent[0]["answers"] == ANSWERS
    assert funnel.unsent(600, now=10**10) == []
    assert "+79995554433" in message.bot.send_message.call_args.args[1]


async def test_someone_elses_contact_is_refused():
    from unittest.mock import AsyncMock

    import bot

    funnel.save_lead(person(12), ANSWERS, "—")
    message = SimpleNamespace(
        contact=SimpleNamespace(user_id=99, phone_number="+70000000000"),
        from_user=SimpleNamespace(id=12),
        answer=AsyncMock(),
    )
    await bot.on_contact(message)
    assert funnel.leads()[0]["phone"] is None
