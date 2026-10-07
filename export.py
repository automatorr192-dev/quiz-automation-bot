import io
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from quiz import QUIZ

MSK = timezone(timedelta(hours=3))


def leads_xlsx(rows: list[dict]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Лиды квиза"
    head = ["Дата", "Имя", "Telegram", "Телефон", "Источник", *[item["q"] for item in QUIZ]]
    sheet.append(head)
    for row in rows:
        sheet.append(
            [
                datetime.fromtimestamp(row["at"], MSK).replace(tzinfo=None),
                row.get("name"),
                f"@{row['username']}" if row.get("username") else row["user_id"],
                row.get("phone"),
                row.get("source"),
                *[row["answers"].get(item["key"]) for item in QUIZ],
            ]
        )
    fill = PatternFill("solid", fgColor="E8EEF9")
    for index, cell in enumerate(sheet[1], start=1):
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[cell.column_letter].width = 17 if index < 6 else 28
    for (cell,) in sheet.iter_rows(min_row=2, max_col=1):
        cell.number_format = "DD.MM.YYYY HH:MM"
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
